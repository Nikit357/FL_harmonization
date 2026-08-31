#!/usr/bin/env python
"""Add metric groups L, M and N to `supplementary_260824/Supplementary File 2.xlsx`.

Groups L (marker gene correlation preservation), M (cross-batch rank agreement) and N
(predictive validation) are the blind-check metrics from
`harmonization-metrics/marker_and_predictive_validation_plan_260819.md`. This script:

* Appends 5 descriptive rows to `Table_S2_metrics` (2 for L, 1 for M, 2 for N), at the same
  level of detail and citation format as the existing A-K rows.
* Appends 82 new scalar-metric rows to `Metric_polarity` with polarity 0 -- one per scalar
  key that `compute_group_l` / `_m` / `_n` in `compute_batch_metrics.py` actually returns
  (dict-valued keys such as `mk_rho_by_gene` are excluded: `run_metrics_concat.py` splits
  those into the three long-format tables, never the wide CSV this sheet mirrors). The
  `xb_rank_agree_{diagnosis}` columns are enumerated from the real
  `Diagnosis_cell_type_unified` levels in `comb_ann_unified.csv`, excluding the one level
  with a single sample (mathematically cannot form a same-biology cross-batch pair).
* Adds 'Metric group' (A-N letter, or 'Metadata' for run-provenance columns produced by no
  metric-group function) and 'Metric type' ('Global'/'Local'/'Distributional similarity'/
  'Other', the vocabulary NAR_review_260727.md §5.3 fixed to match Figures 3 and 9E) columns
  to `Metric_polarity`, classifying all 229 pre-existing rows plus the 82 new ones.

The base-name -> (group, type) map is derived by reading, for every existing base name,
which `compute_group_*` function in `compute_batch_metrics.py` produces it, and matching its
statistical character against the same four Table_S2_metrics type labels already used for
that letter group.

Usage:
    source ~/venvs/collagen_3_11/bin/activate
    python figure_drawing_scripts/add_LMN_metrics_260824.py
"""

import os
import copy
import openpyxl
from openpyxl.styles import Font, Alignment

PATH = os.path.expanduser("~/FL_harmonization/figures_for_article/supplementary_260824/Supplementary File 2.xlsx")

# ── Group L/M/N scalar metric columns (from compute_batch_metrics.py + run_metrics_concat.py) ──

DIAG_LEVELS = [
    "Diffuse_Large_B_Cell_Lymphoma", "Follicular_Lymphoma", "Memory", "Naive", "B_cells",
    "Plasma", "Burkitt_Lymphoma", "Bone_marrow_CD19+", "Centroblast", "Centrocyte", "GC",
    "High_Grade_B_Cell_Lymphoma", "Other", "Extranodal_Marginal_Zone_Lymphoma",
    "Double_Hit_Lymphoma", "MZ", "Immature", "Plasmablast", "Mantle_Cell_Lymphoma",
    # Chronic_Lymphocytic_Leukemia (n=1) excluded: cannot form a same-biology,
    # different-batch pair with only one sample in the cohort.
]

L_BOOKKEEPING = [
    "mk_n_panel_genes_used", "mk_n_genes_null", "mk_panel_coverage_frac",
    "mk_is_self_reference", "mk_n_cohorts_used", "mk_n_cohorts_skipped_small",
]
L_STAT = [
    "mk_rho_mean_all_genes", "mk_rho_median_all_genes", "mk_rho_p10_all_genes",
    "mk_rho_min_all_genes", "mk_rho_frac_genes_above_0.9", "mk_rho_mean_by_cohort_mean",
    "mk_rho_mean_markers", "mk_rho_mean_housekeeping", "mk_rho_marker_minus_hk",
]

M_BOOKKEEPING = [
    "xb_subsampled", "xb_n_samples_used", "xb_n_pairs_same_bio", "xb_n_pairs_diff_bio",
    "xb_n_panel_genes_used",
]
M_STAT = [
    "xb_rank_agree", "xb_rank_agree_median", "xb_rank_disagree_diffbio",
    "xb_rank_disagree_diffbio_median", "xb_rank_agree_ratio",
] + [f"xb_rank_agree_{d}" for d in DIAG_LEVELS]

N_TOP = ["pv_n_perm", "pv_n_genes_dropped_na", "pv_n_samples_dropped_na"]
N_SUFFIXES = [
    "n_folds", "n_folds_multiclass", "f1_macro_mean", "f1_macro_median", "f1_macro_std",
    "f1_macro_min", "f1_weighted_mean", "bal_acc_mean", "mcc_mean", "auc_macro_mean",
    "f1_macro_perm_mean", "f1_macro_perm_std", "auc_macro_perm_mean", "f1_macro_delta",
    "f1_perm_pvalue",
]
N_METRICS = []
for prefix, classes in (("pv_lobo3", ["FL", "DLBCL", "Normal_B"]), ("pv_lobo2", ["FL", "DLBCL"])):
    for suf in N_SUFFIXES:
        N_METRICS.append(f"{prefix}_{suf}")
    for cls in classes:
        N_METRICS.append(f"{prefix}_f1_per_class_{cls}")

NEW_METRICS = []
for m in L_BOOKKEEPING:
    NEW_METRICS.append((m, "L", "Other"))
for m in L_STAT:
    NEW_METRICS.append((m, "L", "Distributional similarity"))
for m in M_BOOKKEEPING:
    NEW_METRICS.append((m, "M", "Other"))
for m in M_STAT:
    NEW_METRICS.append((m, "M", "Distributional similarity"))
for m in N_TOP:
    NEW_METRICS.append((m, "N", "Other"))
for m in N_METRICS:
    NEW_METRICS.append((m, "N", "Other"))

assert len(NEW_METRICS) == len(set(m for m, _, _ in NEW_METRICS)), "duplicate metric name"
print("New L/M/N scalar metrics to add:", len(NEW_METRICS))
print("  L:", len(L_BOOKKEEPING) + len(L_STAT), " M:", len(M_BOOKKEEPING) + len(M_STAT), " N:", len(N_TOP) + len(N_METRICS))

# ── (group, type) classification for the 101 existing base-metric prefixes ──

BATCHY_COLS = {
    "COHORT_LABEL", "PLATFORM_RNA", "RNASEQ_SOURCE", "RNA_BATCH",
    "Diagnosis_cell_type_unified", "Major_group", "TUMOR_NORMAL",
}


def strip_suffix(name: str) -> str:
    for col in sorted(BATCHY_COLS, key=len, reverse=True):
        if name.endswith("_" + col):
            return name[: -(len(col) + 1)]
    return name


# base name -> (group, type)
BASE_MAP = {
    # Group A -- PCA-based variance decomposition
    "r2": ("A", "Global"),
    **{f"r2_pc{i}": ("A", "Global") for i in range(1, 11)},
    "PCReg": ("A", "Global"),
    "dsc": ("A", "Global"),
    "dsc_pvalue": ("A", "Global"),
    # Group B -- neighbor-based integration metrics
    "kbet_acceptance_rate": ("B", "Local"),
    "clisi_mean": ("B", "Local"),
    "ilisi_mean": ("B", "Local"),
    "ilisi_norm": ("B", "Local"),
    "asw_batch": ("B", "Global"),
    "asw_batch_norm": ("B", "Global"),
    "asw_bio": ("B", "Global"),
    "asw_bio_norm": ("B", "Global"),
    "cms_fraction_mixed": ("B", "Global"),
    "cms_mean": ("B", "Global"),
    # Group C -- UMAP/tSNE embedding metrics
    "tsne_centroid_disp": ("C", "Global"),
    "umap_centroid_disp": ("C", "Global"),
    "tsne_entropy_norm": ("C", "Local"),
    "tsne_entropy_mean": ("C", "Local"),
    "umap_entropy_norm": ("C", "Local"),
    "umap_entropy_mean": ("C", "Local"),
    # Group D -- distribution comparison metrics
    "ks_frac_sig": ("D", "Distributional similarity"),
    "ks_mean_D": ("D", "Distributional similarity"),
    "ks_cohort_within_batch_frac_sig": ("D", "Distributional similarity"),
    "ks_cohort_within_batch_mean_D": ("D", "Distributional similarity"),
    "per_gene_batch_mean_cv": ("D", "Distributional similarity"),
    # Group E -- data quality / descriptive statistics
    "n_samples": ("E", "Other"),
    "n_genes": ("E", "Other"),
    "n_batches": ("E", "Other"),
    "n_samples_per_batch_min": ("E", "Other"),
    "n_samples_per_batch_max": ("E", "Other"),
    "n_samples_per_batch_median": ("E", "Other"),
    "n_samples_per_batch_sd": ("E", "Other"),
    "n_cohorts": ("E", "Other"),
    "n_diagnosis_groups": ("E", "Other"),
    "zero_fraction_global": ("E", "Other"),
    "fraction_genes_below_1": ("E", "Other"),
    "zero_fraction_by_batch_max": ("E", "Other"),
    "zero_fraction_by_batch_min": ("E", "Other"),
    "below_1_fraction_by_batch_max": ("E", "Other"),
    "below_1_fraction_by_batch_min": ("E", "Other"),
    "exp_min": ("E", "Other"),
    "exp_max": ("E", "Other"),
    "exp_median": ("E", "Other"),
    "exp_std": ("E", "Other"),
    "exp_p01": ("E", "Other"),
    "exp_p05": ("E", "Other"),
    "exp_p25": ("E", "Other"),
    "exp_p75": ("E", "Other"),
    "exp_p95": ("E", "Other"),
    "exp_p99": ("E", "Other"),
    "per_batch_median_cv": ("E", "Other"),
    "n_cohorts_bimodal": ("E", "Other"),
    "fraction_cohorts_bimodal": ("E", "Other"),
    "n_cohorts_zero_inflated_bimodal": ("E", "Other"),
    "fraction_cohorts_zero_inflated_bimodal": ("E", "Other"),
    # Group G -- graph connectivity
    "graph_connectivity": ("G", "Local"),
    # Group H -- pairwise Euclidean distances
    "dist_ratio": ("H", "Global"),
    "avg_inter_dist": ("H", "Global"),
    "avg_intra_dist": ("H", "Global"),
    # Group I -- WaterMelon score
    "wm": ("I", "Global"),
    "wm_mean_batch": ("I", "Global"),
    "wm_mean_bio": ("I", "Global"),
    "wm_ratio_bio_batch": ("I", "Global"),
    "wm_subsampled": ("I", "Other"),
    # Group J -- per-PC variance explained
    "pct_var_cum_top10": ("J", "Global"),
    **{f"pct_var_pc{i}": ("J", "Global") for i in range(1, 11)},
    # Group K -- NA retention / failure detection
    "pct_genes_noNA": ("K", "Other"),
    "pct_na_cells": ("K", "Other"),
    "pct_samples_noNA": ("K", "Other"),
    "pct_samples_allNA": ("K", "Other"),
    "pct_genes_allNA": ("K", "Other"),
    "n_genes_noNA": ("K", "Other"),
    "n_genes_allNA": ("K", "Other"),
    "n_samples_allNA": ("K", "Other"),
    "n_samples_noNA": ("K", "Other"),
    "n_na_cells": ("K", "Other"),
    # Run/benchmark provenance -- not produced by any A-N metric-group function
    "compute_time_s": ("Metadata", "Other"),
    "harshness_level": ("Metadata", "Other"),
    "is_best": ("Metadata", "Other"),
    "is_shambhala": ("Metadata", "Other"),
    "shambhala_p_key": ("Metadata", "Other"),
    "shambhala_q_key": ("Metadata", "Other"),
}
for m, g, t in NEW_METRICS:
    BASE_MAP[m] = (g, t)


def classify(name: str):
    base = strip_suffix(name)
    if base in BASE_MAP:
        return BASE_MAP[base]
    if name in BASE_MAP:
        return BASE_MAP[name]
    return ("?", "?")


# ── apply to workbook ──

wb = openpyxl.load_workbook(PATH, data_only=False)

# --- Table_S2_metrics: 5 new descriptive rows for groups L, M, N ---
ws2 = wb["Table_S2_metrics"]

KOTLOV = (
    "Kotlov,N., Bagaev,A., Revuelta,M.V., Phillip,J.M., Cacciapuoti,M.T., Antysheva,Z., "
    "Svekolkin,V., Tikhonova,E., Miheecheva,N., Kuzkina,N., et al. (2021) Clinical and "
    "Biological Subtypes of B-cell Lymphoma Revealed by Microenvironmental Signatures. "
    "Cancer Discov., 11, 1468–1489."
)
HOLMES = (
    "Holmes,A.B., Corinaldesi,C., Shen,Q., Kumar,R., Compagno,N., Wang,Z., Nitzan,M., "
    "Grunstein,E., Pasqualucci,L., Dalla-Favera,R. and Basso,K. (2020) Single-cell analysis "
    "of germinal-center B cells informs on lymphoma cell of origin and outcome. J. Exp. Med., "
    "217, e20200483."
)
PANEL_SOURCE = f"This study (custom implementation); Spearman rank correlation. Marker panel from {KOTLOV} and {HOLMES}"
CHICCO = (
    "Chicco,D. and Jurman,G. (2020) The advantages of the Matthews correlation coefficient "
    "(MCC) over F1 score and accuracy in binary classification evaluation. BMC Genomics, 21, 6."
)
OJALA = (
    "Ojala,M. and Garriga,G.C. (2010) Permutation tests for studying classifier performance. "
    "J. Mach. Learn. Res., 11, 1833–1863."
)

NEW_S2_ROWS = [
    (
        "Marker gene expression profile preservation (Spearman ρ)",
        "Distributional similarity",
        "L",
        "For every gene in the marker/housekeeping panel and every cohort with ≥20 samples, "
        "the Spearman correlation between that gene's expression across the cohort's samples "
        "before harmonization (the matched 01_raw reference) and after. Reported as the grand "
        "mean, median, 10th percentile, minimum and fraction of genes with ρ > 0.9 over all "
        "panel genes, and as the mean over cohorts. Any harmonizer applying a per-batch monotone "
        "transform of each gene leaves within-cohort ranks unchanged and scores ~1.0 here by "
        "construction, so a near-1.0 value is not by itself evidence of quality.",
        PANEL_SOURCE,
    ),
    (
        "Marker-gene-versus-housekeeping-gene correlation contrast",
        "Distributional similarity",
        "L",
        "The mean before/after Spearman ρ over biologically variable marker genes minus the "
        "mean ρ over 15 housekeeping genes. Housekeeping genes carry little biological "
        "signal, so a harmonizer that genuinely reshuffles samples should drive their ρ "
        "toward 0 while marker ρ stays high; a per-batch monotone transform, however, "
        "preserves ranks on both panels equally and this contrast reads as a relative check "
        "rather than an absolute floor.",
        PANEL_SOURCE,
    ),
    (
        "Cross-batch same- versus different-biology rank agreement",
        "Distributional similarity",
        "M",
        "Within one harmonized matrix, panel genes are rank-transformed within each sample and "
        "the Spearman correlation is computed for every pair of samples drawn from different "
        "RNA_BATCH levels. Pairs are split into same-biology (same Diagnosis_cell_type_unified) "
        "and different-biology; the margin between the two means is the biology-specific signal "
        "that harmonization should increase — a method that merely collapses every sample "
        "toward a common profile raises both, leaving the margin flat. Computed from the job's "
        "own matrix only; the 01_raw attempt supplies the pre-harmonization baseline by "
        "construction.",
        PANEL_SOURCE,
    ),
    (
        "Leave-one-batch-out diagnosis prediction (F1, AUC, balanced accuracy, MCC)",
        "Other",
        "N",
        "A logistic regression on the first 10 principal components (fitted on the training "
        "batches only, held-out batch projected) is evaluated in a leave-one-RNA_BATCH-out "
        "scheme for a 3-class target (FL / DLBCL / Normal_B) and a 2-class target (FL vs DLBCL). "
        "Macro/weighted F1, balanced accuracy, Matthews correlation coefficient and one-vs-rest "
        "AUC are reported as the mean over folds, together with per-class F1 and the fold counts "
        "used for each mean.",
        f"This study (custom implementation); leave-one-batch-out cross-validated multinomial "
        f"logistic regression. {CHICCO}",
    ),
    (
        "Label-permutation negative control for diagnosis prediction",
        "Other",
        "N",
        "The class labels are shuffled globally 100 times (fixed seed) and the leave-one-batch-out "
        "evaluation (both the 3-class and 2-class targets) is repeated on each permutation, "
        "reusing the label-independent per-fold PCA projections. Reports the permuted mean and "
        "standard deviation of macro-F1 and AUC, the observed-minus-permuted delta, and an "
        "empirical p-value (fraction of permutations reaching or exceeding the observed macro-F1, "
        "with +1 added to numerator and denominator so it is never exactly zero).",
        f"This study (custom implementation); label-permutation empirical null for classifier "
        f"performance. {OJALA}",
    ),
]

data_font = copy.copy(ws2["A2"].font)
data_align = copy.copy(ws2["A2"].alignment)
start_row = ws2.max_row + 1
for i, row in enumerate(NEW_S2_ROWS):
    r = start_row + i
    for c, val in enumerate(row, start=1):
        cell = ws2.cell(row=r, column=c, value=val)
        cell.font = data_font
        cell.alignment = data_align
print(f"Table_S2_metrics: appended {len(NEW_S2_ROWS)} rows (now {ws2.max_row})")

# --- Metric_polarity: new 'Metric group' / 'Metric type' columns + new L/M/N rows ---
ws1 = wb["Metric_polarity"]
header_font = copy.copy(ws1["A1"].font)
data_font = copy.copy(ws1["A2"].font)
data_align = copy.copy(ws1["A2"].alignment)
n_existing = ws1.max_row  # includes header at row 1

ws1.cell(row=1, column=3, value="Metric group").font = header_font
ws1.cell(row=1, column=4, value="Metric type").font = header_font

unresolved = []
for r in range(2, n_existing + 1):
    name = ws1.cell(row=r, column=1).value
    group, mtype = classify(name)
    if group == "?":
        unresolved.append(name)
    c3 = ws1.cell(row=r, column=3, value=group)
    c4 = ws1.cell(row=r, column=4, value=mtype)
    c3.font = data_font
    c3.alignment = data_align
    c4.font = data_font
    c4.alignment = data_align

if unresolved:
    print("UNRESOLVED existing metric names (need manual mapping):", unresolved)
else:
    print(f"Classified all {n_existing - 1} existing metrics into group/type")

next_row = n_existing + 1
for name, group, mtype in NEW_METRICS:
    ws1.cell(row=next_row, column=1, value=name).font = data_font
    ws1.cell(row=next_row, column=1).alignment = data_align
    c2 = ws1.cell(row=next_row, column=2, value=0)
    c2.font = data_font
    c2.alignment = data_align
    c3 = ws1.cell(row=next_row, column=3, value=group)
    c3.font = data_font
    c3.alignment = data_align
    c4 = ws1.cell(row=next_row, column=4, value=mtype)
    c4.font = data_font
    c4.alignment = data_align
    next_row += 1

ws1.column_dimensions["C"].width = 14.71
ws1.column_dimensions["D"].width = 20.71
print(f"Metric_polarity: now {ws1.max_row} rows (was {n_existing})")

wb.save(PATH)
print("Saved:", PATH)
