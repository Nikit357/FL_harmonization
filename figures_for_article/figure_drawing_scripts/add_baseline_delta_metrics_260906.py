#!/usr/bin/env python
"""Add the unharmonized-baseline contrast family to `Metric_polarity`.

Groups M and N are single-matrix metrics: unlike Group L they carry no internal
before/after comparison, so the "did harmonization help?" question is answered in the
analysis layer instead, by looking up the `01_raw` attempt of the same
(strategy, imputation) pair. `figures_helpers.attach_raw_baseline()` — and cell 6 of
both `correlation_prediction_metrics_analysis*.ipynb` notebooks, which inline the same
logic — emit a `{metric}_raw` and a `{metric}_delta` column for seven M/N metrics, plus
the bookkeeping flags that qualify them. None of those 18 columns exist in
`metrics_comprehensive.csv`, so the earlier passes over this workbook, which enumerated
the sheet from the compute pipeline's own output, never saw them.

This script appends:

* 14 rows — `{m}_raw` / `{m}_delta` for the seven metrics in `BASELINE_METRICS`
  (`xb_rank_agree`, `xb_rank_agree_ratio`, `xb_rank_disagree_diffbio`,
  `pv_lobo{3,2}_f1_macro_mean`, `pv_lobo{3,2}_auc_macro_mean`), polarity 0, in the
  group and type of their parent metric.
* 4 rows — `baseline_missing`, `is_raw`, `group`, `pv_perm_comparable`, as `Metadata`,
  matching how `is_best` and `harshness_level` are already recorded.
* 1 `Table_S2_metrics` row describing the construction.

**Name collision to be aware of.** `pv_lobo3_f1_macro_delta` (already in the sheet) is
produced by `compute_group_n()` and is *observed minus the label-permutation null*.
`pv_lobo3_f1_macro_mean_delta` (added here) is *observed minus the unharmonized
baseline*. Same prefix, different comparison; both explanations say so explicitly.

Usage:
    source ~/venvs/collagen_3_11/bin/activate
    python figure_drawing_scripts/add_baseline_delta_metrics_260906.py
"""

import copy
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent.parent          # figures_for_article/
PATH = HERE / "supplementary_260824" / "Supplementary File 2.xlsx"

BASELINE_METHOD = "01_raw"
HELPER = "figures_helpers.attach_raw_baseline()"

# Parent metric -> (group, type, short description used in the generated text).
BASELINE_METRICS: list[tuple[str, str, str, str]] = [
    ("xb_rank_agree", "M", "Distributional similarity",
     "the mean cross-batch rank agreement between same-biology sample pairs"),
    ("xb_rank_agree_ratio", "M", "Distributional similarity",
     "the same-biology minus different-biology cross-batch agreement margin"),
    ("xb_rank_disagree_diffbio", "M", "Distributional similarity",
     "the mean cross-batch rank agreement between different-biology sample pairs"),
    ("pv_lobo3_f1_macro_mean", "N", "Other",
     "the leave-one-batch-out macro-F1 of the 3-class diagnosis classifier"),
    ("pv_lobo2_f1_macro_mean", "N", "Other",
     "the leave-one-batch-out macro-F1 of the 2-class (FL vs DLBCL) classifier"),
    ("pv_lobo3_auc_macro_mean", "N", "Other",
     "the leave-one-batch-out one-vs-rest AUC of the 3-class diagnosis classifier"),
    ("pv_lobo2_auc_macro_mean", "N", "Other",
     "the leave-one-batch-out AUC of the 2-class (FL vs DLBCL) classifier"),
]

WHY_EXTERNAL = (
    "Groups M and N are computed from one matrix and therefore carry no internal "
    "before/after comparison the way Group L does; the unharmonized value is supplied "
    "by the 01_raw attempt, which is itself a row of the benchmark. The lookup is done "
    "once in the analysis layer rather than inside every job, where it would repeat the "
    "same number about sixty times per (strategy, imputation) pair."
)
POST1_CAVEAT = (
    "The baseline is always the post-removal-off 01_raw attempt, so on a post_rm = True "
    "row the two sides rest on different sample sets — post-removal drops "
    "method-specific outlier batches. These metrics are means over sample pairs and "
    "folds rather than sample-matched statistics, so the comparison still holds, but "
    "post1 deltas are approximate and post0 / post1 must be reported separately."
)
NOT_PIPELINE = (
    "Derived during the downstream article analysis, not by compute_batch_metrics.py: "
    "it is absent from metrics_comprehensive.csv and is created on load by "
    f"{HELPER}, whose logic cell 6 of "
    "correlation_prediction_metrics_analysis.ipynb and its _narrow_set variant inline."
)

NEW_ROWS: list[tuple[str, str, str, str, str]] = []   # metric, group, type, expl, calc

for parent, group, mtype, what in BASELINE_METRICS:
    collision = ""
    if parent.endswith("_f1_macro_mean"):
        sibling = parent.replace("_f1_macro_mean", "_f1_macro_delta")
        collision = (
            f" Do not confuse this with {sibling}, a different column that the metrics "
            f"pipeline itself produces: that one is the observed macro-F1 minus the "
            f"LABEL-PERMUTATION null, whereas this one is the observed value minus the "
            f"UNHARMONIZED baseline."
        )
    NEW_ROWS.append((
        f"{parent}_raw", group, mtype,
        f"Group {group} baseline column: the value of {parent} — {what} — measured on "
        f"the unharmonized {BASELINE_METHOD} attempt of this row's own "
        f"(strategy, imputation) pair. {WHY_EXTERNAL} {NOT_PIPELINE}",
        f"baseline = the rows with method == '{BASELINE_METHOD}' and post_rm == False, "
        f"indexed by (strat, imp); {parent}_raw = baseline[{parent}] reindexed onto this "
        f"row's (strat, imp) pair. A row whose pair has no {BASELINE_METHOD} attempt "
        f"gets NaN and is flagged by baseline_missing. Implemented in {HELPER}.",
    ))
    if parent == "xb_rank_agree_ratio":
        headline = (
            " This is the headline of the whole baseline-contrast family and the one "
            "the best-versus-rest figures plot as 'margin gain vs raw': it is the only "
            "member that cannot be inflated by a method which merely collapses every "
            "sample toward one common profile, because such a method raises "
            "xb_rank_agree and xb_rank_disagree_diffbio together and leaves this "
            "margin's delta near zero."
        )
    else:
        headline = (
            " Positive means harmonization improved the metric relative to leaving the "
            "matrix unharmonized; near zero means it changed nothing that this metric "
            "can see."
        )
    NEW_ROWS.append((
        f"{parent}_delta", group, mtype,
        f"Group {group} baseline contrast: {parent} minus its unharmonized value, i.e. "
        f"how far harmonization moved {what} away from the {BASELINE_METHOD} matrix of "
        f"the same (strategy, imputation) pair.{headline}{collision} {NOT_PIPELINE}",
        f"{parent}_delta = {parent} - {parent}_raw. {POST1_CAVEAT} Implemented in "
        f"{HELPER}.",
    ))

NEW_ROWS += [
    (
        "baseline_missing", "Metadata", "Other",
        "Metadata/provenance, assigned during downstream article analysis: a boolean "
        "flag marking that this attempt's (strategy, imputation) pair has no "
        f"{BASELINE_METHOD} attempt with post-removal off, so every *_raw and *_delta "
        "column on the row is NaN. Read it before quoting any delta: a missing baseline "
        "is not a delta of zero.",
        "pair_index = MultiIndex of (strat, imp) over every row; baseline_missing = "
        f"~pair_index.isin(lookup.index), where lookup holds the {BASELINE_METHOD}, "
        f"post_rm == False rows indexed the same way. Implemented in {HELPER}.",
    ),
    (
        "is_raw", "Metadata", "Other",
        "Metadata/provenance, assigned during downstream article analysis: a boolean "
        f"flag marking that this row IS an unharmonized {BASELINE_METHOD} attempt. Used "
        "to drop the baselines from best-versus-rest comparisons and from the noise "
        "comparison of §3b, where a row whose delta is zero by construction would "
        "distort the spread.",
        f"df_lmn['is_raw'] = df_lmn['method'] == '{BASELINE_METHOD}'. Set in cell 5 of "
        "correlation_prediction_metrics_analysis.ipynb and its _narrow_set variant.",
    ),
    (
        "group", "Metadata", "Other",
        "Metadata/provenance, assigned during downstream article analysis: the "
        "best-versus-rest label every §2 panel splits on — 'Best (clustermap)' for the "
        "15 curated top approaches, 'Rest' for every other attempt. The string form of "
        "is_best, kept as its own column so it can be used directly as a seaborn "
        "grouping variable.",
        "df_lmn['group'] = np.where(df_lmn['is_best'], 'Best (clustermap)', 'Rest'); "
        "is_best itself comes from the curated (method, imp, strat) list documented on "
        "the is_best row. Set in cell 5 of correlation_prediction_metrics_analysis.ipynb "
        "and its _narrow_set variant.",
    ),
    (
        "pv_perm_comparable", "Metadata", "Other",
        "Metadata/provenance, assigned during downstream article analysis: a boolean "
        "flag marking that this row's permutation count equals the run's expected one, "
        "so its empirical p-value shares the same floor as the rest of the table. The "
        "floor is 1/(n_perm + 1), so a row computed with a different pv_n_perm cannot be "
        "compared against the others and is excluded from the p-value panel.",
        "PV_N_PERM_EXPECTED = 20 (the 260824 run; floor 0.0476). "
        "df_lmn['pv_perm_comparable'] = df_lmn['pv_n_perm'] == PV_N_PERM_EXPECTED, and "
        "False everywhere if the pv_n_perm column is absent. Set in cell 5 of "
        "correlation_prediction_metrics_analysis.ipynb and its _narrow_set variant.",
    ),
]

assert len(NEW_ROWS) == 18, len(NEW_ROWS)
assert len({m for m, *_ in NEW_ROWS}) == 18, "duplicate metric name"

NEW_S2_ROW = (
    "Unharmonized-baseline contrast for the single-matrix metrics (Δ versus 01_raw)",
    "Distributional similarity",
    "M, N",
    "Group L compares a harmonized matrix against its unharmonized reference inside the "
    "metric itself. Groups M and N cannot: both are computed from a single matrix, so a "
    "raw M or N value says how good an attempt is, not how much harmonization changed "
    "it. The contrast is therefore formed in the analysis layer, where the 01_raw "
    "attempt of the same (strategy, imputation) pair — itself a row of the benchmark — "
    "supplies the 'before' value. Seven M/N metrics gain a `_raw` column (the baseline "
    "value) and a `_delta` column (this attempt minus the baseline): the three Group M "
    "rank-agreement metrics and the macro-F1 and AUC of both Group N targets. "
    "`xb_rank_agree_ratio_delta` is the headline of the family, because it is the only "
    "one that cannot be inflated by a method that collapses every sample toward a "
    "common profile. " + POST1_CAVEAT + " A row whose pair has no baseline is flagged "
    "by `baseline_missing` rather than silently scored as zero.",
    "This study (custom implementation); paired contrast against the unharmonized "
    "attempt of the same strategy and imputation. Implemented in "
    "figures_for_article/figure_drawing_scripts/figures_helpers.py, "
    "attach_raw_baseline(); the same logic is inlined in cell 6 of "
    "harmonization-metrics/correlation_prediction_metrics_analysis.ipynb and its "
    "_narrow_set variant.",
)


def main() -> None:
    wb = openpyxl.load_workbook(PATH)

    ws1 = wb["Metric_polarity"]
    # Rewrite in place when the row is already there, so a wording fix can be re-applied
    # to a workbook an earlier run of this script already touched.
    where = {ws1.cell(row=r, column=1).value: r for r in range(2, ws1.max_row + 1)}
    font = copy.copy(ws1["A2"].font)
    align = copy.copy(ws1["A2"].alignment)
    n_add = n_upd = 0
    for name, group, mtype, expl, calc in NEW_ROWS:
        if name in where:
            row, n_upd = where[name], n_upd + 1
        else:
            row, n_add = ws1.max_row + 1, n_add + 1
        for col, val in enumerate((name, 0, group, mtype, expl, calc), start=1):
            cell = ws1.cell(row=row, column=col, value=val)
            cell.font = font
            cell.alignment = align
    print(f"Metric_polarity: {n_add} rows appended, {n_upd} rewritten "
          f"(now {ws1.max_row})")

    ws2 = wb["Table_S2_metrics"]
    existing_s2 = {ws2.cell(row=r, column=1).value: r
                   for r in range(2, ws2.max_row + 1)}
    row = existing_s2.get(NEW_S2_ROW[0], ws2.max_row + 1)
    font = copy.copy(ws2["A2"].font)
    align = copy.copy(ws2["A2"].alignment)
    for col, val in enumerate(NEW_S2_ROW, start=1):
        cell = ws2.cell(row=row, column=col, value=val)
        cell.font = font
        cell.alignment = align
    print(f"Table_S2_metrics: row written at {row} (now {ws2.max_row})")

    wb.save(PATH)
    print("Saved:", PATH)


if __name__ == "__main__":
    main()
