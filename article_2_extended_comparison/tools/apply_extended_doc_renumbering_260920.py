#!/usr/bin/env python3
"""Renumber the extended document's figures to Article 2's scheme, as tracked changes.

Article 2 promotes Extended Figures 1 to 3 to main Figures 1 to 3 and demotes Extended
Figures 4 to 13 to Supplementary Figures 1 to 10, and the Figma frames are renamed to
match. The extended document still calls them Extended Figures, so text and canvas would
disagree. This script writes a **new** snapshot with the renumbering applied; v2 is
never overwritten, because the audit cites it by name.

**The collision this script exists to handle.** The extended document also cites the
source benchmark's own figures: "Figure 3" seven times, "Figure 4" once, and
Supplementary Figures 6, 7 and 17. Renaming its Extended Figure 3 to "Figure 3" would
put two different meanings on one label. Daniil's decision of 2026-09-20 is to qualify
the borrowed references instead, so every one of them gains "of the source benchmark"
**before** any Extended Figure is renamed. Run the two passes in that order or the
qualification pass will also catch the newly renamed figures.

The renumbering itself follows the longest-match discipline of
`figures_for_article/CLAUDE.md`: two-digit numbers before one-digit ones, and the
"Extended Figure 7 and 8" span before plain "Extended Figure 7". Each replacement is
grown with surrounding context until it is unique inside its paragraph, so no occurrence
can be matched by the wrong key.

    source ~/venvs/collagen_3_11/bin/activate
    cd article_2_extended_comparison
    python tools/apply_extended_doc_renumbering_260920.py
"""
from __future__ import annotations

import argparse
import re
import shutil
import sys
from pathlib import Path

SKILLS = Path(__file__).resolve().parents[2] / ".claude" / "skills"
sys.path.insert(0, str(SKILLS))

import docx  # noqa: E402
from nar_review_tools import safe_tracked_replace  # noqa: E402
import word_rewrite_trackchanges as wt  # noqa: E402

AUTHOR = "Claude (Article 2)"
DATE = "2026-09-20T00:00:00Z"
QUALIFIER = " of the source benchmark"

# Pass 1: the figures that belong to the source benchmark, not to this document.
# The lookbehinds keep "Extended Figure 3" and "Supplementary Figure 3" out of the plain
# "Figure 3" key; the lookahead stops an already-qualified reference matching again.
SOURCE_REFS = re.compile(
    r"(?<!Extended )(?<!Supplementary )\bFigure (?:3|4)\b(?!" + QUALIFIER + r")"
    r"|\bSupplementary Figure (?:6|7|17)\b(?!" + QUALIFIER + r")")

# Pass 2: this document's own figures, longest key first.
RENUMBER = [
    ("Extended Figures 1-3", "Figures 1-3"),
    ("Extended Figures 1–3", "Figures 1–3"),
    ("Extended Figure 7 and 8", "Supplementary Figure 4 and 5"),
    ("Extended Figure 10", "Supplementary Figure 7"),
    ("Extended Figure 11", "Supplementary Figure 8"),
    ("Extended Figure 12", "Supplementary Figure 9"),
    ("Extended Figure 13", "Supplementary Figure 10"),
    ("Extended Figure 1", "Figure 1"),
    ("Extended Figure 2", "Figure 2"),
    ("Extended Figure 3", "Figure 3"),
    ("Extended Figure 4", "Supplementary Figure 1"),
    ("Extended Figure 5", "Supplementary Figure 2"),
    ("Extended Figure 6", "Supplementary Figure 3"),
    ("Extended Figure 7", "Supplementary Figure 4"),
    ("Extended Figure 8", "Supplementary Figure 5"),
    ("Extended Figure 9", "Supplementary Figure 6"),
]

MAX_CONTEXT = 120


def paragraph_text(p) -> str:
    """The text safe_tracked_replace itself sees: direct-child runs only."""
    return "".join("".join(t.text or "" for t in r.findall(wt.qn("w:t")))
                   for r in p if r.tag == wt.qn("w:r"))


def unique_span(text: str, start: int, end: int, replacement: str):
    """Grow the match outwards until the literal is unique in the paragraph.

    safe_tracked_replace edits the first occurrence of the literal it is given, so a key
    that appears twice in one paragraph would silently edit the wrong one.
    """
    for pad in range(0, MAX_CONTEXT + 1, 10):
        a, b = max(0, start - pad), min(len(text), end + pad)
        old = text[a:b]
        if text.count(old) == 1:
            return old, text[a:start] + replacement + text[end:b]
    return None, None


def apply_pass(paragraph, find_next) -> int:
    """Apply one replacement at a time until the paragraph has no match left."""
    n = 0
    while True:
        text = paragraph_text(paragraph)
        found = find_next(text)
        if not found:
            return n
        start, end, replacement = found
        old, new = unique_span(text, start, end, replacement)
        if old is None:
            print(f"  ! could not make {text[start:end]!r} unique in its paragraph",
                  file=sys.stderr)
            return n
        (_, status), = safe_tracked_replace(paragraph, [(old, new)])
        if status != "ok":
            print(f"  ! {status}: {text[start:end]!r}", file=sys.stderr)
            return n
        n += 1


def next_source_ref(text: str):
    m = SOURCE_REFS.search(text)
    return (m.start(), m.end(), m.group(0) + QUALIFIER) if m else None


def next_renumber(text: str):
    best = None
    for old, new in RENUMBER:
        i = text.find(old)
        # Longest key wins at the same position, which is what RENUMBER's order gives;
        # the earliest position wins overall so the paragraph is rewritten left to right.
        if i >= 0 and (best is None or i < best[0]):
            best = (i, i + len(old), new)
    return best


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--src",
                    default="manuscript_versions/"
                            "Harmonization_metrics_extended_260802_v2.docx")
    ap.add_argument("--dst",
                    default="manuscript_versions/"
                            "Harmonization_metrics_extended_260920_v3.docx")
    args = ap.parse_args()

    src, dst = Path(args.src), Path(args.dst)
    if dst.exists():
        print(f"{dst.name} already exists; delete it first if you mean to rebuild it",
              file=sys.stderr)
        return 1
    shutil.copy2(src, dst)

    wt.set_revision_identity(AUTHOR, DATE)
    document = docx.Document(str(dst))

    qualified = sum(apply_pass(p._p, next_source_ref) for p in document.paragraphs)
    renumbered = sum(apply_pass(p._p, next_renumber) for p in document.paragraphs)

    document.save(str(dst))
    print(f"  {qualified} source-benchmark references qualified")
    print(f"  {renumbered} Extended Figure references renumbered")
    print(f"  wrote {dst.name} ({qualified + renumbered} tracked changes, "
          f"author {AUTHOR})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
