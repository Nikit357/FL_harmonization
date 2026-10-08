#!/usr/bin/env python
"""Phase 5 deliverable: the revision as Word tracked changes against Daniil's edited copy.

    python tools/build_tracked_revision_260924.py

Base   `manuscript/FL_metric_classes_F1000_260917_edited_Daniil.docx` (never modified)
Edits  `workflow_runs/260924_run1/03_revision_edits_final.json` (assembled + renumbered)
Output `manuscript/FL_metric_classes_F1000_260925.docx`

How each record becomes tracked changes (author "Claude (Article 2 revision 2026-09-25)"):

* `edit` — the paragraph is diffed against its base text word by word and only the
  changed spans become `w:del` + `w:ins`. Mendeley citations (`w:sdt`) are atoms that the
  diff can never touch, so all 46 content controls survive byte for byte (plan §C1; this
  is why `word_rewrite_trackchanges.tracked_replace`, which rebuilds the paragraph, is not
  used). A span inside one of Daniil's own tracked insertions is deleted by a `w:del`
  nested inside his `w:ins`, the OOXML form of deleting another author's insertion.
* `delete` — every visible run becomes a tracked deletion, including runs inside Daniil's
  insertions (which the stock `delete_paragraph` helper would leave visible), and the
  paragraph mark is marked deleted; a table cell's row is marked deleted too.
* `insert` — a new paragraph after the current position, every run tracked-inserted;
  `**bold**` spans become bold runs and a Markdown table becomes a tracked-inserted Word
  table in the style of the manuscript's existing table.

Checks, written to `workflow_runs/260924_run1/22_docx_build_report.md`:

1. Accept All: every paragraph reads exactly as the revised text.
2. Rejecting Claude's revisions (and only those) restores Daniil's document, paragraph by
   paragraph, with his own 18 revisions still in place.
3. 46 Mendeley content controls with unchanged contents; the numeric citations intact.
4. Orphan paragraphs left holding only citations after a deletion.
"""
from __future__ import annotations

import copy
import difflib
import json
import re
import sys
import zipfile
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path.home() / ".claude/skills"))
import word_rewrite_trackchanges as wr  # noqa: E402

BASE = ROOT / "manuscript/FL_metric_classes_F1000_260917_edited_Daniil.docx"
OUT = ROOT / "manuscript/FL_metric_classes_F1000_260925.docx"
EDITS = ROOT / "workflow_runs/260924_run1/03_revision_edits_final.json"
REPORT = ROOT / "workflow_runs/260924_run1/22_docx_build_report.md"
AUTHOR = "Claude (Article 2 revision 2026-09-25)"
DATE = "2026-09-25T00:00:00Z"

W_P, W_R, W_T, W_TAB = qn("w:p"), qn("w:r"), qn("w:t"), qn("w:tab")
W_INS, W_DEL, W_DT, W_SDT = qn("w:ins"), qn("w:del"), qn("w:delText"), qn("w:sdt")
W_PPR, W_RPR = qn("w:pPr"), qn("w:rPr")
CIT = ""  # one character per Mendeley content control in the diff strings
SAFE_KIDS = {W_RPR, W_T, qn("w:lastRenderedPageBreak")}


def is_mendeley(el) -> bool:
    pr = el.find(qn("w:sdtPr"))
    return el.tag == W_SDT and pr is not None and b"MENDELEY" in etree.tostring(pr)


def is_claude(el) -> bool:
    return el.get(qn("w:author")) == AUTHOR


# ── visible text units ────────────────────────────────────────────────────────

def units(p) -> list[dict]:
    """The visible text of a paragraph as a list of units in document order.

    kind `run`: a direct run we may split; `drun`: a run inside Daniil's w:ins (split
    inside his insertion); `cit`: a Mendeley content control (one CIT character, never
    edited); `fixed`: visible text we must not split (a run with a tab, a field, or our
    own insertions). Runs inside a w:del are invisible and skipped.
    """
    out = []
    for ch in p:
        if ch.tag == W_R:
            txt = "".join(t.text or "" for t in ch.findall(W_T))
            plain = all(k.tag in SAFE_KIDS for k in ch)
            if txt or ch.find(W_TAB) is not None:
                if ch.find(W_TAB) is not None:
                    txt = "".join("\t" if k.tag == W_TAB else (k.text or "")
                                  for k in ch if k.tag in (W_T, W_TAB))
                out.append({"kind": "run" if plain else "fixed", "el": ch, "text": txt})
        elif ch.tag == W_SDT:
            out.append({"kind": "cit" if is_mendeley(ch) else "fixed", "el": ch,
                        "text": CIT if is_mendeley(ch) else wr.text_of(ch)})
        elif ch.tag == W_INS:
            for r in ch.findall(W_R):
                txt = "".join(t.text or "" for t in r.findall(W_T))
                if txt:
                    plain = all(k.tag in SAFE_KIDS for k in r)
                    kind = "fixed" if is_claude(ch) or not plain else "drun"
                    out.append({"kind": kind, "el": r, "text": txt, "ins": ch})
        # w:del (invisible), bookmarks, comment ranges, proofErr: no visible text
    return out


def visible(p) -> str:
    return "".join(u["text"] for u in units(p))


# ── the minimal edit ─────────────────────────────────────────────────────────

def _split_run(run, cut: int):
    """Split a plain run at character `cut`; return (left, right), either may be None."""
    txt = "".join(t.text or "" for t in run.findall(W_T))
    if cut <= 0:
        return None, run
    if cut >= len(txt):
        return run, None
    base = run.find(W_RPR)
    left, right = wr.make_run(txt[:cut], base=base), wr.make_run(txt[cut:], base=base)
    run.addprevious(left)
    run.addprevious(right)
    run.getparent().remove(run)
    return left, right


def _delete_run(run):
    """Turn a whole plain run into a tracked deletion, in place; return the w:del."""
    base = run.find(W_RPR)
    txt = "".join(t.text or "" for t in run.findall(W_T))
    d = wr.wrap_del(wr.make_del_run(txt, base=base))
    run.addprevious(d)
    run.getparent().remove(run)
    return d


def _ins_run(text: str, base):
    return wr.wrap_ins(wr.make_run(text, base=base))


def apply_edit(p, start: int, end: int, new: str) -> None:
    """Replace visible characters [start, end) of paragraph p by `new`, tracked."""
    us, pos = units(p), 0
    spans = []
    for u in us:
        a, b = pos, pos + len(u["text"])
        spans.append((a, b, u))
        pos = b
    covered = [(a, b, u) for a, b, u in spans if a < end and b > start and end > start]
    for a, b, u in covered:
        if u["kind"] in ("cit", "fixed"):
            raise ValueError(f"edit [{start},{end}) touches a {u['kind']} unit "
                             f"{u['text'][:30]!r}")
    anchor, base_rpr = None, None
    for a, b, u in covered:  # left to right: split, then delete the covered middle
        run = u["el"]
        base_rpr = run.find(W_RPR)
        lo, hi = max(start, a) - a, min(end, b) - a
        left, rest = _split_run(run, lo)
        mid, _right = _split_run(rest, hi - lo)
        d = _delete_run(mid)
        anchor = d
    if new:
        if anchor is None:  # a pure insertion at `start`
            inside = [(a, b, u) for a, b, u in spans if a < start < b]
            ending = [(a, b, u) for a, b, u in spans if b == start and b > a]
            starting = [(a, b, u) for a, b, u in spans if a == start and b > a]
            if inside:
                a, b, u = inside[0]
                if u["kind"] not in ("run", "drun"):
                    raise ValueError(f"insertion inside a {u['kind']} unit")
                left, _ = _split_run(u["el"], start - a)
                anchor = left
            elif ending:
                anchor = ending[-1][2]["el"]
            elif starting:
                u = starting[0][2]
                ins = _ins_run(new, u["el"].find(W_RPR) if u["kind"] in ("run", "drun")
                               else None)
                target = u["ins"] if u["kind"] == "drun" else u["el"]
                target.addprevious(ins)
                return
            else:  # empty paragraph: append
                p.append(_ins_run(new, None))
                return
            base_rpr = anchor.find(W_RPR) if anchor.tag == W_R else None
        ins = _ins_run(new, base_rpr)
        parent_ins = next((a for a in anchor.iterancestors() if a.tag == W_INS), None)
        if parent_ins is not None and not is_claude(parent_ins):
            # Inside Daniil's insertion: an ins cannot nest, so split his w:ins after the
            # anchor and put ours between the two halves.
            tail = [s for s in anchor.itersiblings()]
            parent_ins.addnext(ins)
            if tail:
                second = copy.deepcopy(parent_ins)
                for c in list(second):
                    second.remove(c)
                second.set(qn("w:id"), wr.nid())
                for s in tail:
                    second.append(s)
                ins.addnext(second)
        elif anchor.getparent() is not None and anchor.getparent().tag == W_INS \
                and not is_claude(anchor.getparent()):
            anchor.getparent().addnext(ins)
        else:
            anchor.addnext(ins)


TOKEN_RE = re.compile(r"\s+|\w+|[^\w\s]")


def diff_edits(old: str, new: str) -> list[tuple[int, int, str]]:
    """Word-level changes turning `old` into `new`, as (start, end, replacement)."""
    a, b = TOKEN_RE.findall(old), TOKEN_RE.findall(new)
    sm = difflib.SequenceMatcher(None, a, b, autojunk=False)
    offs = [0]
    for t in a:
        offs.append(offs[-1] + len(t))
    raw = [(offs[i1], offs[i2], "".join(b[j1:j2]))
           for tag, i1, i2, j1, j2 in sm.get_opcodes() if tag != "equal"]
    # Merge changes separated only by a single whitespace token: one readable
    # replacement instead of a run of one-word edits.
    merged = []
    for s, e, r in raw:
        if merged and s - merged[-1][1] <= 1 and old[merged[-1][1]:s].strip() == "":
            ps, pe, pr = merged[-1]
            merged[-1] = (ps, e, pr + old[pe:s] + r)
        else:
            merged.append((s, e, r))
    return merged


def edit_paragraph(p, new_text: str) -> int:
    old = visible(p)
    edits = diff_edits(old, new_text)
    # A change may not swallow a citation: split it at every CIT on both sides.
    final = []
    for s, e, r in edits:
        seg_old, n_cit = old[s:e], old[s:e].count(CIT)
        if n_cit == 0:
            final.append((s, e, r))
            continue
        if r.count(CIT) != n_cit:
            raise ValueError(f"a change moves a citation: {seg_old!r} -> {r!r}")
        o_parts, n_parts = seg_old.split(CIT), r.split(CIT)
        pos = s
        for op, np_ in zip(o_parts, n_parts):
            if op != np_:
                final.append((pos, pos + len(op), np_))
            pos += len(op) + 1
    for s, e, r in sorted(final, key=lambda x: x[0], reverse=True):
        apply_edit(p, s, e, r)
    return len(final)


# ── deletion and insertion ────────────────────────────────────────────────────

def delete_whole_paragraph(p) -> None:
    ppr = p.find(W_PPR)
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        p.insert(0, ppr)
    rpr = ppr.find(W_RPR)
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        ppr.append(rpr)
    rpr.append(wr._rev_attrs(OxmlElement("w:del")))
    for ch in list(p):
        if ch.tag == W_R and (ch.findall(W_T) or ch.find(W_TAB) is not None):
            for t in ch.findall(W_T):
                t.tag = W_DT
            d = wr._rev_attrs(OxmlElement("w:del"))
            ch.addprevious(d)
            d.append(ch)
        elif ch.tag == W_INS and not is_claude(ch):
            for r in list(ch.findall(W_R)):
                for t in r.findall(W_T):
                    t.tag = W_DT
                d = wr._rev_attrs(OxmlElement("w:del"))
                r.addprevious(d)
                d.append(r)
        elif ch.tag == W_SDT and is_mendeley(ch):
            raise ValueError("refusing to delete a paragraph that holds a citation")
    tr = next((a for a in p.iterancestors() if a.tag == qn("w:tr")), None)
    if tr is not None:
        trpr = tr.find(qn("w:trPr"))
        if trpr is None:
            trpr = OxmlElement("w:trPr")
            tr.insert(0, trpr)
        if trpr.find(W_DEL) is None:
            trpr.append(wr._rev_attrs(OxmlElement("w:del")))


def md_runs(text: str, base=None) -> list:
    """Runs of an inserted paragraph; `**bold**` spans become bold runs."""
    runs = []
    for i, part in enumerate(re.split(r"\*\*", text)):
        if part:
            runs.append(wr.make_run(part, bold=bool(i % 2), base=base))
    return runs


def ins_para(text: str, style: str):
    p = OxmlElement("w:p")
    p.append(wr._ppr(style, "ins"))
    runs = md_runs(text)
    if runs:
        p.append(wr.wrap_ins(*runs))
    return p


def ins_table(rows: list[list[str]], tbl_template):
    """A tracked-inserted Word table: every row carries w:ins in its trPr."""
    tbl = OxmlElement("w:tbl")
    if tbl_template is not None and tbl_template.find(qn("w:tblPr")) is not None:
        tbl.append(copy.deepcopy(tbl_template.find(qn("w:tblPr"))))
    grid = OxmlElement("w:tblGrid")
    for _ in rows[0]:
        grid.append(OxmlElement("w:gridCol"))
    tbl.append(grid)
    for ri, cells in enumerate(rows):
        tr = OxmlElement("w:tr")
        trpr = OxmlElement("w:trPr")
        trpr.append(wr._rev_attrs(OxmlElement("w:ins")))
        tr.append(trpr)
        for c in cells:
            tc = OxmlElement("w:tc")
            p = ins_para(f"**{c}**" if ri == 0 else c, "Compact")
            tc.append(p)
            tr.append(tc)
        tbl.append(tr)
    return tbl


def split_md_table(text: str):
    """(paragraphs before, table rows or None, paragraphs after) of an insert."""
    lines = text.splitlines()
    tl = [i for i, ln in enumerate(lines) if ln.strip().startswith("|")]
    if not tl:
        return [ln for ln in re.split(r"\n\s*\n", text) if ln.strip()], None, []
    before = "\n".join(lines[:tl[0]]).strip()
    after = "\n".join(lines[tl[-1] + 1:]).strip()
    rows = []
    for ln in lines[tl[0]:tl[-1] + 1]:
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
            continue
        rows.append(cells)
    return ([before] if before else []), rows, ([after] if after else [])


def clean(text: str) -> str:
    """The revised text as it must read in Word: provenance comments stripped, citation
    tokens rendered as CIT characters, runs of spaces collapsed."""
    t = re.sub(r"<!--.*?-->", "", text or "", flags=re.S)
    t = re.sub(r"⟦CIT:.*?⟧", CIT, t)
    t = re.sub(r"(?<=\S) {2,}(?=\S)", " ", t).strip()
    return t


# ── views for validation ─────────────────────────────────────────────────────

def view(p, reject_claude: bool, accept_all: bool) -> str | None:
    """Paragraph text under a review state. None when the paragraph itself disappears.

    reject_claude: Claude's ins removed, Claude's del restored, Daniil's revisions shown
    as [+x+] / [-x-] so they are compared too. accept_all: every ins kept, every del
    dropped, citations as CIT characters.
    """
    ppr = p.find(W_PPR)
    mark = ppr.find(W_RPR) if ppr is not None else None
    if mark is not None:
        m_ins, m_del = mark.find(W_INS), mark.find(W_DEL)
        if reject_claude and m_ins is not None and is_claude(m_ins):
            return None
        if accept_all and m_del is not None:
            return None
    out = []

    def walk(node, ctx):
        for ch in node:
            if ch.tag == W_SDT and is_mendeley(ch):
                out.append(CIT if accept_all else "<CIT:" + wr.text_of(ch) + ">")
            elif ch.tag == W_INS:
                if reject_claude and is_claude(ch):
                    continue
                if reject_claude:
                    out.append("[+")
                    walk(ch, ctx)
                    out.append("+]")
                else:
                    walk(ch, ctx)
            elif ch.tag == W_DEL:
                if accept_all:
                    continue
                if is_claude(ch):
                    walk(ch, "restore")
                else:
                    out.append("[-")
                    walk(ch, "restore")
                    out.append("-]")
            elif ch.tag == W_T:
                out.append(ch.text or "")
            elif ch.tag == W_DT:
                if ctx == "restore" or not accept_all:
                    out.append(ch.text or "")
            elif ch.tag == W_TAB:
                out.append("\t")
            elif ch.tag in (W_R, qn("w:hyperlink"), qn("w:smartTag")):
                walk(ch, ctx)

    walk(p, None)
    return "".join(out)


def word_reject_all(body) -> str:
    """Body text after Word's Reject All: every insertion removed with everything inside
    it (a deletion nested in an insertion goes with it), every other deletion restored.

    The stock `nar_review_tools validate` counts text deleted inside an insertion as
    restored, which Word does not do; this is the check that matches Word.
    """
    out = []

    def walk(node, in_ins):
        for ch in node:
            if ch.tag == W_INS:
                continue
            if ch.tag == W_T:
                out.append(ch.text or "")
            elif ch.tag == W_DT:
                out.append(ch.text or "")
            elif ch.tag == W_TAB:
                out.append("\t")
            else:
                walk(ch, in_ins)

    for p in body.iter(W_P):
        ppr = p.find(W_PPR)
        mark = ppr.find(W_RPR) if ppr is not None else None
        if mark is not None and mark.find(W_INS) is not None:
            continue
        if any(a.tag == qn("w:tr") and a.find(qn("w:trPr")) is not None
               and a.find(qn("w:trPr")).find(W_INS) is not None for a in p.iterancestors()):
            continue
        walk(p, False)
        out.append("\n")
    return "".join(out)


def body_paragraphs(body):
    return [p for p in body.iter(W_P)]


def main() -> int:
    wr.set_revision_identity(AUTHOR, DATE)
    d = docx.Document(str(BASE))
    body = d.element.body
    ids = [int(e.get(qn("w:id"))) for e in body.iter() if e.get(qn("w:id"), "").isdigit()]
    wr._rid[0] = max(ids + [0]) + 10000  # revision ids must not collide with Daniil's
    base_view = {id(p): view(p, reject_claude=True, accept_all=False)
                 for p in body_paragraphs(body)}
    base_order = [view(p, True, False) for p in body_paragraphs(body)]

    # The index numbers the non-empty paragraphs exactly as phase3_paragraph_index.py
    # does, so P<n> is the n-th paragraph with visible text.
    paras = [p for p in body_paragraphs(body) if visible(p).strip()]
    records = json.loads(EDITS.read_text())
    table_tpl = body.find(qn("w:tbl"))
    stats = {"edit": 0, "delete": 0, "insert": 0, "keep": 0, "changes": 0}
    expected = []  # (element, expected accept-all text)
    cursor = None
    problems = []
    for r in records:
        if r["action"] == "insert":
            before, rows, after = split_md_table(r["new"])
            style = "BodyText"
            new_els = []
            for t in before:
                txt = clean(t)
                if txt:
                    new_els.append(ins_para(txt, style))
                    expected.append((new_els[-1], txt.replace("**", "")))
            if rows:
                tbl = ins_table([[clean(c) for c in row] for row in rows], table_tpl)
                new_els.append(tbl)
            for t in after:
                txt = clean(t)
                if txt:
                    new_els.append(ins_para(txt, style))
                    expected.append((new_els[-1], txt.replace("**", "")))
            host = cursor
            while host is not None and host.getparent() is not body:
                host = host.getparent()  # never insert inside a table cell
            wr.insert_after(host, new_els)
            cursor = new_els[-1] if new_els else cursor
            stats["insert"] += 1
            continue
        n = int(r["id"][1:])
        p = paras[n]
        cursor = p
        if r["action"] == "keep":
            stats["keep"] += 1
            expected.append((p, visible(p)))
        elif r["action"] == "delete":
            delete_whole_paragraph(p)
            stats["delete"] += 1
        else:
            new_text = clean(r["new"]).replace("**", "")
            try:
                stats["changes"] += edit_paragraph(p, new_text)
            except (ValueError, AttributeError) as exc:
                problems.append(f"{r['id']}: {type(exc).__name__}: {exc}")
            expected.append((p, new_text))
            stats["edit"] += 1

    # ── checks ──
    accept_bad = [(wr.text_of(el)[:60], exp[:60]) for el, exp in expected
                  if el.tag == W_P and (view(el, False, True) or "") != exp]
    # Splitting one of Daniil's insertions around ours leaves two adjacent insertions of
    # his once ours is rejected -- the same author, date and text, so the boundary
    # marker between them carries no content and is dropped before comparing.
    norm = lambda v: v.replace("+][+", "").replace("-][-", "")
    rejected = [norm(v) for v in (view(p, True, False) for p in body_paragraphs(body))
                if v is not None]
    base_order = [norm(v) for v in base_order]
    reject_ok = rejected == base_order
    first_bad = next((i for i, (a, b) in enumerate(zip(rejected, base_order)) if a != b),
                     None)
    d.save(str(OUT))
    xml = zipfile.ZipFile(OUT).read("word/document.xml").decode("utf-8", "ignore")
    base_xml = zipfile.ZipFile(BASE).read("word/document.xml").decode("utf-8", "ignore")
    sdt_base = [wr.text_of(s) for s in docx.Document(str(BASE)).element.body.iter(W_SDT)]
    sdt_new = [wr.text_of(s) for s in docx.Document(str(OUT)).element.body.iter(W_SDT)]
    acc = "\n".join(view(p, False, True) or "" for p in
                    docx.Document(str(OUT)).element.body.iter(W_P))
    numeric = {c: acc.count(c) for c in ("[38]", "[50,51]", "[52]")}
    word_reject_ok = (word_reject_all(docx.Document(str(OUT)).element.body)
                      == word_reject_all(docx.Document(str(BASE)).element.body))
    orphans = sum(1 for p in docx.Document(str(OUT)).element.body.iter(W_P)
                  if (v := view(p, False, True)) is not None and v.strip()
                  and set(v.replace(CIT, "").strip(" ;,()")) == set())
    claude_revs = sum(1 for e in docx.Document(str(OUT)).element.body.iter()
                      if e.tag in (W_INS, W_DEL) and is_claude(e))
    lines = ["# 22 — Tracked-changes build report (Phase 5), 2026-09-25", "",
             f"Output `{OUT.relative_to(ROOT)}` against `{BASE.relative_to(ROOT)}` "
             f"(unchanged). Author of every new revision: \"{AUTHOR}\".", "",
             f"- records: {stats['keep']} kept, {stats['edit']} edited "
             f"({stats['changes']} tracked changes), {stats['delete']} deleted, "
             f"{stats['insert']} inserted; Claude revision marks: {claude_revs}",
             f"- **Accept All = revised text:** {len(accept_bad) == 0} "
             f"({len(expected) - len(accept_bad)} of {len(expected)} paragraphs)",
             f"- **Reject Claude's revisions = Daniil's copy:** {reject_ok} "
             f"({len(rejected)} vs {len(base_order)} paragraphs"
             + (f"; first difference at paragraph {first_bad}" if first_bad is not None
                else "") + ")",
             f"- **Word's Reject All gives the same text for both files:** {word_reject_ok} "
             "(the stock validator's reject-all differs where Claude deleted text inside "
             "one of Daniil's insertions; Word discards such a nested deletion with the "
             "insertion)",
             f"- **Mendeley content controls:** {xml.count('MENDELEY_CITATION')} "
             f"(base {base_xml.count('MENDELEY_CITATION')}); contents identical: "
             f"{sdt_base == sdt_new}",
             f"- **numeric citations after Accept All:** {numeric} "
             "(expected [38] ×2, [50,51] ×1, [52] ×1)",
             f"- **orphan citation-only paragraphs:** {orphans}",
             f"- edit problems: {len(problems)}"] + [f"  - {x}" for x in problems]
    if accept_bad:
        lines += ["", "## Accept-All mismatches (first 10)", ""] + [
            f"- got `{a}` / expected `{b}`" for a, b in accept_bad[:10]]
    REPORT.write_text("\n".join(lines) + "\n")
    print("\n".join(lines))
    ok = (not accept_bad and reject_ok and word_reject_ok and sdt_base == sdt_new and not problems
          and numeric == {"[38]": 2, "[50,51]": 1, "[52]": 1})
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
